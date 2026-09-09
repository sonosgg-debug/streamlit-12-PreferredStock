"""
app.py
한국증시 우선주 목록 및 투자 지표 비교 대시보드 웹 애플리케이션.
보통주-우선주 매핑, 13개 필수 지표 산출, 엑셀 다운로드, 인터랙티브 상세 차트 제공.
"""

import sys
# Python 3.12+ 및 Streamlit Cloud(Python 3.14) 환경에서 pykrx의 pkg_resources 모듈 임포트 에러 방지용 shim
try:
    import pkg_resources
except Exception:
    try:
        import setuptools.command
        import pkg_resources
    except Exception:
        import types
        pkg_mock = types.ModuleType("pkg_resources")
        pkg_mock.resource_filename = lambda *args, **kwargs: ""
        pkg_mock.resource_string = lambda *args, **kwargs: b""
        pkg_mock.Requirement = type("Requirement", (), {"parse": lambda s: s})
        pkg_mock.get_distribution = lambda *args, **kwargs: type("Dist", (), {"version": "1.0.0"})()
        sys.modules["pkg_resources"] = pkg_mock

import io
import os
import textwrap
import datetime
import streamlit as st
import pandas as pd
import numpy as np
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import data_loader

# 1. 페이지 설정
st.set_page_config(
    page_title="한국증시 우선주 목록 및 투자 지표 비교",
    layout="wide",
    initial_sidebar_state="expanded"
)

# 2. 커스텀 CSS (00 Bookmarks 다크 테마 및 스타일 적용)
st.markdown("""
<style>
    /* Main Background */
    .stApp {
        background-color: #0f172a;
        color: #f8fafc;
        font-family: -apple-system, BlinkMacSystemFont, "Apple SD Gothic Neo", "Malgun Gothic", "맑은 고딕", sans-serif;
    }
    
    /* Main Content Area */
    .main .block-container,
    [data-testid="stMainBlockContainer"] {
        padding-top: 2.5rem !important;
        padding-bottom: 3.5rem !important;
        max-width: 98% !important;
    }
    
    /* Sidebar Styling */
    section[data-testid="stSidebar"] {
        background-color: #1e293b !important;
        border-right: 1px solid #334155;
    }
    
    section[data-testid="stSidebar"] h1, 
    section[data-testid="stSidebar"] h2, 
    section[data-testid="stSidebar"] h3 {
        color: #f8fafc !important;
    }

    /* Input & Select Box styling */
    .stTextInput input, .stSelectbox select {
        background-color: #334155 !important;
        color: #f8fafc !important;
        border: 1px solid #475569 !important;
        border-radius: 6px !important;
    }

    /* Metric Card Styling */
    .metric-card {
        background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%);
        border: 1px solid #334155;
        border-radius: 10px;
        padding: 16px 20px;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.2);
        margin-bottom: 15px;
    }
    .metric-label {
        font-size: 0.85rem;
        color: #94a3b8;
        font-weight: 600;
        margin-bottom: 6px;
    }
    .metric-value {
        font-size: 1.6rem;
        font-weight: 800;
        color: #f8fafc;
    }
    .metric-highlight {
        color: #38bdf8;
    }
    .metric-green {
        color: #34d399;
    }
    .metric-amber {
        color: #fbbf24;
    }

    /* Section Subheaders */
    .section-header {
        font-size: 1.25rem;
        font-weight: 700;
        color: #e2e8f0;
        margin-top: 10px;
        margin-bottom: 12px;
        display: flex;
        align-items: center;
        gap: 8px;
    }

    /* Button Styling */
    .stButton button[kind="primary"] {
        background-color: #2563eb !important;
        color: #ffffff !important;
        border: none !important;
        font-weight: 600 !important;
        border-radius: 6px !important;
        transition: all 0.2s ease !important;
    }
    .stButton button[kind="primary"]:hover {
        background-color: #1d4ed8 !important;
        box-shadow: 0 0 10px rgba(37, 99, 235, 0.4) !important;
    }

    /* Download button */
    .stDownloadButton button {
        background-color: #059669 !important;
        color: #ffffff !important;
        border: none !important;
        font-weight: 600 !important;
        border-radius: 6px !important;
        transition: all 0.2s ease !important;
    }
    .stDownloadButton button:hover {
        background-color: #047857 !important;
        box-shadow: 0 0 10px rgba(5, 150, 105, 0.4) !important;
    }
</style>
""", unsafe_allow_html=True)


# 3. 데이터 로딩 캐시 함수
@st.cache_data(ttl=3600, show_spinner=False)
def get_cached_market_data(force_refresh=False):
    return data_loader.load_market_data(force_refresh=force_refresh)


# 4. 엑셀 파일 생성 헬퍼 함수
def create_excel_download(df_export):
    """
    스타일 및 서식이 적용된 엑셀 파일 바이너리 생성
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # 데이터 시트 작성
        sheet_name = "우선주_투자지표_비교"
        df_export.to_excel(writer, sheet_name=sheet_name, index=False)
        
        workbook = writer.book
        worksheet = writer.sheets[sheet_name]

        # 헤더 스타일
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # 데이터 셀 서식
        data_font = Font(name="맑은 고딕", size=10)
        border_thin = Side(style='thin', color="CBD5E1")
        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        # 헤더 서식 적용
        for col_idx in range(1, len(df_export.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = cell_border

        # 데이터 행 서식 적용
        for row_idx in range(2, len(df_export) + 2):
            for col_idx, col_name in enumerate(df_export.columns, start=1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = cell_border

                # 컬럼 유형별 정렬 및 숫자 서식
                if col_name in ["시장", "보통주명", "우선주명"]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_name in ["보통주가", "우선주가", "보통주 배당금", "우선주 배당금"]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '#,##0'
                elif col_name in ["시가총액(보통주)"]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '#,##0'
                elif col_name in ["배당성향", "괴리율(%)", "보통주 배당수익률(A)", "우선주 배당수익률(B)"]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '0.00'
                elif col_name in ["배당수익률 비율(B/A)"]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '0.00'
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # 자동 열 너비 계산
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val.encode('euc-kr', errors='ignore')))
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        worksheet.row_dimensions[1].height = 28

    return output.getvalue()


# 5. 세션 상태 초기화
if 'market_selection' not in st.session_state:
    st.session_state.market_selection = "전체"
if 'selected_stock_idx' not in st.session_state:
    st.session_state.selected_stock_idx = 0
if 'force_reload' not in st.session_state:
    st.session_state.force_reload = False


# ==========================================
# 6. 왼쪽 사이드 패널 (사이드바)
# ==========================================
with st.sidebar:
    st.markdown("<h2 style='color: #8AB4F8; font-size: 1.3rem; margin-top: 0;'>⚙️ 검색 및 필터</h2>", unsafe_allow_html=True)
    
    # 1) 시장 선택 (디폴트: "전체")
    market_choice = st.radio(
        "시장 선택",
        options=["전체", "코스피", "코스닥"],
        index=0,
        horizontal=True,
        help="조회할 주식 시장을 선택합니다."
    )
    
    # 2) 조회 버튼
    btn_search = st.button("🔍 조회", type="primary", use_container_width=True)
    if btn_search:
        st.session_state.market_selection = market_choice
        st.rerun()

    st.markdown("<hr style='border: 0; height: 1px; background-color: #334155; margin: 18px 0;'>", unsafe_allow_html=True)
    
    # 3) 스마트 필터
    st.markdown("<div style='font-size: 0.95rem; font-weight: 700; color: #cbd5e1; margin-bottom: 8px;'>🎯 스마트 필터</div>", unsafe_allow_html=True)
    
    search_keyword = st.text_input("종목명 검색", placeholder="예: 삼성, 현대, LG...", help="보통주 또는 우선주 이름으로 검색합니다.")
    
    col_f1, col_f2 = st.columns(2)
    with col_f1:
        min_div_yield = st.number_input("최소 배당률(%)", min_value=0.0, max_value=20.0, value=0.0, step=0.5, help="0.0%는 전체 종목 표시")
    with col_f2:
        min_discount = st.number_input(
            "최소 괴리율(%)",
            min_value=-1000.0,
            max_value=90.0,
            value=0.0,
            step=5.0,
            help="기본값 0.0% (보통주보다 저렴한 종목만 표시). 마이너스 입력 시 역괴리 종목도 포함"
        )

    only_dividend_paying = st.checkbox("배당 지급 종목만 보기", value=False, help="우선주 배당금이 0원 초과인 종목만 필터링합니다.")

    st.markdown("<hr style='border: 0; height: 1px; background-color: #334155; margin: 18px 0;'>", unsafe_allow_html=True)

    # 4) 캐시 갱신 버튼
    if st.button("🔄 최신 데이터 강제 갱신", use_container_width=True):
        st.cache_data.clear()
        st.session_state.force_reload = True
        st.rerun()


# ==========================================
# 7. 데이터 로드 및 필터링
# ==========================================
force_refresh = st.session_state.force_reload
st.session_state.force_reload = False

with st.spinner("KRX 시장 및 펀더멘털 데이터를 불러오는 중입니다..."):
    df_raw, target_date = get_cached_market_data(force_refresh=force_refresh)

if df_raw.empty:
    st.error("데이터를 불러올 수 없습니다. 인터넷 연결 및 KRX 인증 설정을 확인하세요.")
    st.stop()

# 시장 필터 적용
df_filtered = df_raw.copy()
if market_choice == "코스피":
    df_filtered = df_filtered[df_filtered["시장"] == "코스피"]
elif market_choice == "코스닥":
    df_filtered = df_filtered[df_filtered["시장"] == "코스닥"]

# 스마트 필터 적용
if search_keyword.strip():
    kw = search_keyword.strip().lower()
    df_filtered = df_filtered[
        df_filtered["보통주명"].str.lower().str.contains(kw) |
        df_filtered["우선주명"].str.lower().str.contains(kw)
    ]

if min_div_yield > 0:
    df_filtered = df_filtered[df_filtered["우선주 배당수익률(B)"] >= min_div_yield]

# 최소 괴리율(%) 기본값(0.0%) 적용: 보통주 대비 할인 거래되는 종목 필터링
df_filtered = df_filtered[df_filtered["괴리율(%)"] >= min_discount]

if only_dividend_paying:
    df_filtered = df_filtered[df_filtered["우선주 배당금"] > 0]


# ==========================================
# 8. 메인 영역: 타이틀 영역
# ==========================================
# 타이틀 색상은 00 Bookmarks 폴더 앱과 완벽히 동일한 #8AB4F8, 서브타이틀 없음
st.markdown(
    "<h1 style='color: #8AB4F8 !important; font-weight: 800; font-size: 1.95rem; margin-top: 0; margin-bottom: 0.3rem; letter-spacing: -0.5px;'>"
    "한국증시 우선주 목록 및 투자 지표 비교"
    "</h1>",
    unsafe_allow_html=True
)

# 기준일자 메타 정보 표시
date_formatted = f"{target_date[:4]}-{target_date[4:6]}-{target_date[6:]}"
st.markdown(
    f"<div style='font-size: 0.85rem; color: #94a3b8; margin-bottom: 12px;'>"
    f"기준일: <span style='color: #38bdf8; font-weight: 600;'>{date_formatted}</span> (전일 종가 기준) &nbsp;|&nbsp; 제공처: <span style='color: #cbd5e1;'>KRX 정보데이터시스템</span>"
    f"</div>",
    unsafe_allow_html=True
)

# [가로 선 1]: 타이틀 영역과 데이터 영역 사이
st.markdown("<hr style='border: 0; height: 1px; background-color: #334155; margin: 10px 0 20px 0;'>", unsafe_allow_html=True)


# ==========================================
# 9. 메인 영역: 핵심 요약 KPI 지표 카드
# ==========================================
col_kpi1, col_kpi2, col_kpi3, col_kpi4 = st.columns(4)

total_count = len(df_filtered)
avg_discount = df_filtered["괴리율(%)"].mean() if total_count > 0 else 0.0
avg_pref_yield = df_filtered["우선주 배당수익률(B)"].mean() if total_count > 0 else 0.0

# 최고 괴리율 종목
if total_count > 0:
    top_discount_row = df_filtered.loc[df_filtered["괴리율(%)"].idxmax()]
    top_discount_info = f"{top_discount_row['우선주명']} ({top_discount_row['괴리율(%)']:.1f}%)"
else:
    top_discount_info = "-"

with col_kpi1:
    with st.container(border=True):
        st.metric("📊 조회된 종목 수", f"{total_count:,} 개")

with col_kpi2:
    with st.container(border=True):
        st.metric("📉 평균 괴리율", f"{avg_discount:.2f} %")

with col_kpi3:
    with st.container(border=True):
        st.metric("💰 평균 우선주 배당률", f"{avg_pref_yield:.2f} %")

with col_kpi4:
    with st.container(border=True):
        st.metric("🏆 최대 괴리율 종목", top_discount_info)


# ==========================================
# 10. 메인 영역: 조회 결과 데이터 영역
# ==========================================
col_title, col_dl = st.columns([8, 2])
with col_title:
    st.markdown("<div class='section-header'>📋 종목별 비교 데이터 테이블</div>", unsafe_allow_html=True)
with col_dl:
    # 사용자 요청 13개 항목 컬럼
    export_cols = [
        "시장",
        "보통주명",
        "우선주명",
        "시가총액(보통주)",
        "배당성향",
        "보통주가",
        "우선주가",
        "괴리율(%)",
        "보통주 배당금",
        "보통주 배당수익률(A)",
        "우선주 배당금",
        "우선주 배당수익률(B)",
        "배당수익률 비율(B/A)"
    ]
    df_export = df_filtered[export_cols].copy()
    
    excel_bytes = create_excel_download(df_export)
    file_name = f"한국증시_우선주_투자지표_비교_{target_date}.xlsx"
    
    st.download_button(
        label="📥 엑셀 다운로드 (.xlsx)",
        data=excel_bytes,
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

# 화면 표시용 데이터프레임 가공 (시가총액 억원 단위 및 컬럼 서식)
df_display = df_filtered.copy()

# 시가총액(보통주)를 억원 단위로 변환하여 보기 쉽게 구성
df_display["시가총액(보통주)_억"] = (df_display["시가총액(보통주)"] / 1e8).round(0)

# 표시할 컬럼 지정 (원래 순서 준수)
display_cols = [
    "시장",
    "보통주명",
    "우선주명",
    "시가총액(보통주)_억",
    "배당성향",
    "보통주가",
    "우선주가",
    "괴리율(%)",
    "보통주 배당금",
    "보통주 배당수익률(A)",
    "우선주 배당금",
    "우선주 배당수익률(B)",
    "배당수익률 비율(B/A)"
]

# 컬럼 서식 설정 (전체 열 너비를 헤더 글자가 가려지지 않는 선에서 콤팩트하게 축소)
column_config = {
    "시장": st.column_config.TextColumn("시장", width="small"),
    "보통주명": st.column_config.TextColumn("보통주명", width="small"),
    "우선주명": st.column_config.TextColumn("우선주명", width="small"),
    "시가총액(보통주)_억": st.column_config.NumberColumn("시가총액(보통주)", format="%,.0f 억", width="small"),
    "배당성향": st.column_config.NumberColumn("배당성향", format="%.2f%%", width="small"),
    "보통주가": st.column_config.NumberColumn("보통주가", format="%,.0f 원", width="small"),
    "우선주가": st.column_config.NumberColumn("우선주가", format="%,.0f 원", width="small"),
    "괴리율(%)": st.column_config.NumberColumn("괴리율(%)", format="%.2f%%", width="small"),
    "보통주 배당금": st.column_config.NumberColumn("보통주 배당금", format="%,.0f 원", width="small"),
    "보통주 배당수익률(A)": st.column_config.NumberColumn("보통주 배당수익률(A)", format="%.2f%%", width="small"),
    "우선주 배당금": st.column_config.NumberColumn("우선주 배당금", format="%,.0f 원", width="small"),
    "우선주 배당수익률(B)": st.column_config.NumberColumn("우선주 배당수익률(B)", format="%.2f%%", width="small"),
    "배당수익률 비율(B/A)": st.column_config.NumberColumn("배당수익률 비율(B/A)", format="%.2f 배", width="small"),
}

# st.dataframe으로 엑셀 스타일 테이블 및 오름차순/내림차순 인터랙티브 정렬 지원
selection = st.dataframe(
    df_display[display_cols],
    use_container_width=True,
    height=420,
    hide_index=True,
    column_config=column_config,
    on_select="rerun",
    selection_mode="single-row"
)

# 테이블에서 선택된 행이 있으면 해당 종목을 상세 분석 대상으로 지정
selected_stock_name = None
if selection and selection.get("rows"):
    sel_idx = selection["rows"][0]
    if sel_idx < len(df_filtered):
        selected_stock_name = df_filtered.iloc[sel_idx]["우선주명"]


# ==========================================
# 11. [가로 선 2]: 데이터 영역과 상세 영역 사이
# ==========================================
st.markdown("<hr style='border: 0; height: 1px; background-color: #334155; margin: 25px 0 25px 0;'>", unsafe_allow_html=True)


# ==========================================
# 12. 메인 영역: 종목별 상세 정보 영역
# ==========================================
st.markdown("<div class='section-header'>📈 종목별 상세 비교 및 시각화 차트</div>", unsafe_allow_html=True)

# 종목 선택 컨트롤러 (테이블 행 선택과 드롭다운 상호 연동)
stock_options = df_filtered["우선주명"].tolist()

if not stock_options:
    st.info("조건에 맞는 종목이 없습니다.")
    st.stop()

# 디폴트 인덱스 계산
if selected_stock_name and selected_stock_name in stock_options:
    default_stock_index = stock_options.index(selected_stock_name)
else:
    # 디폴트는 삼성전자우 또는 첫 번째 종목
    default_stock_index = stock_options.index("삼성전자우") if "삼성전자우" in stock_options else 0

col_sel1, col_sel2, col_sel3 = st.columns([4, 4, 4])
with col_sel1:
    chosen_pref_name = st.selectbox(
        "분석할 우선주 종목 선택",
        options=stock_options,
        index=default_stock_index,
        help="테이블에서 행을 클릭하거나 목록에서 종목을 직접 선택할 수 있습니다."
    )

target_row = df_filtered[df_filtered["우선주명"] == chosen_pref_name].iloc[0]
chosen_pref_ticker = target_row["우선주코드"]
chosen_com_ticker = target_row["보통주코드"]
chosen_com_name = target_row["보통주명"]

with col_sel2:
    period_label = st.radio(
        "조회 기간",
        options=["3개월", "6개월", "1년", "3년"],
        index=2,
        horizontal=True
    )
    period_map = {"3개월": 3, "6개월": 6, "1년": 12, "3년": 36}
    chosen_months = period_map[period_label]

with col_sel3:
    chart_view_mode = st.radio(
        "주가 비교 기준",
        options=["실제 주가(원)", "수익률(기준일=100)"],
        index=1,
        horizontal=True,
        help="시작일 기준 상대 수익률 또는 실제 주가를 비교합니다."
    )

# 상세 요약 카드
with st.container(border=True):
    col_m1, col_m2, col_m3, col_m4, col_m5, col_m6 = st.columns(6)
    with col_m1:
        st.metric("보통주 종목", f"{chosen_com_name}", help=f"티커: {chosen_com_ticker}")
    with col_m2:
        st.metric("우선주 종목", f"{chosen_pref_name}", help=f"티커: {chosen_pref_ticker}")
    with col_m3:
        st.metric("현재 괴리율", f"{target_row['괴리율(%)']:.2f}%")
    with col_m4:
        st.metric("보통주 배당률", f"{target_row['보통주 배당수익률(A)']:.2f}%")
    with col_m5:
        st.metric("우선주 배당률", f"{target_row['우선주 배당수익률(B)']:.2f}%")
    with col_m6:
        ratio_str = f"{target_row['배당수익률 비율(B/A)']:.2f}배" if pd.notna(target_row['배당수익률 비율(B/A)']) else "-"
        st.metric("배당수익률 비율(B/A)", ratio_str)

# 시계열 데이터 로드
with st.spinner("시계열 주가 및 괴리율 데이터를 로드하는 중..."):
    df_hist = data_loader.load_price_history(chosen_pref_ticker, chosen_com_ticker, months=chosen_months)

if df_hist.empty:
    st.warning("선택한 종목의 시계열 주가 데이터를 불러올 수 없습니다.")
else:
    # 2단 레이아웃으로 차트 배치
    col_chart1, col_chart2 = st.columns([6, 6])

    # [차트 1]: 보통주 vs 우선주 주가 추이 (또는 정규화 수익률)
    with col_chart1:
        fig1 = go.Figure()
        
        if chart_view_mode == "수익률(기준일=100)":
            fig1.add_trace(go.Scatter(
                x=df_hist.index,
                y=df_hist['보통주_정규화'],
                mode='lines',
                name=f"{chosen_com_name} (보통주)",
                line=dict(color='#60a5fa', width=2)
            ))
            fig1.add_trace(go.Scatter(
                x=df_hist.index,
                y=df_hist['우선주_정규화'],
                mode='lines',
                name=f"{chosen_pref_name} (우선주)",
                line=dict(color='#f472b6', width=2.5)
            ))
            y_title = "상대 수익률 (시작일=100)"
        else:
            fig1.add_trace(go.Scatter(
                x=df_hist.index,
                y=df_hist['보통주가'],
                mode='lines',
                name=f"{chosen_com_name} (보통주)",
                line=dict(color='#60a5fa', width=2)
            ))
            fig1.add_trace(go.Scatter(
                x=df_hist.index,
                y=df_hist['우선주가'],
                mode='lines',
                name=f"{chosen_pref_name} (우선주)",
                line=dict(color='#f472b6', width=2.5)
            ))
            y_title = "주가 (원)"

        fig1.update_layout(
            title=dict(
                text=f"<b>{chosen_com_name} vs {chosen_pref_name} 주가 추이</b>",
                font=dict(color="#f8fafc", size=15)
            ),
            template="plotly_dark",
            paper_bgcolor="#1e293b",
            plot_bgcolor="#0f172a",
            hovermode="x unified",
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
            margin=dict(l=40, r=20, t=50, b=40),
            yaxis=dict(title=y_title, gridcolor="#334155"),
            xaxis=dict(gridcolor="#334155")
        )
        st.plotly_chart(fig1, use_container_width=True)

    # [차트 2]: 역사적 괴리율(%) 추이 및 평균 밴드
    with col_chart2:
        fig2 = go.Figure()

        # 일별 괴리율
        fig2.add_trace(go.Scatter(
            x=df_hist.index,
            y=df_hist['괴리율'],
            mode='lines',
            name="일별 괴리율(%)",
            line=dict(color='#34d399', width=1.5),
            fill='tozeroy',
            fillcolor='rgba(52, 211, 153, 0.08)'
        ))

        # 20일 이동평균선
        fig2.add_trace(go.Scatter(
            x=df_hist.index,
            y=df_hist['괴리율_MA20'],
            mode='lines',
            name="20일 이동평균",
            line=dict(color='#38bdf8', width=2, dash='dot')
        ))

        # 기간 평균선
        hist_avg_discount = df_hist['괴리율'].mean()
        fig2.add_hline(
            y=hist_avg_discount,
            line_dash="dash",
            line_color="#fbbf24",
            annotation_text=f"기간 평균 ({hist_avg_discount:.1f}%)",
            annotation_position="top right",
            annotation_font=dict(color="#fbbf24", size=11)
        )

        fig2.update_layout(
            title=dict(
                text=f"<b>괴리율 추이 및 평균 밴드 ({period_label})</b>",
                font=dict(color="#f8fafc", size=15)
            ),
            template="plotly_dark",
            paper_bgcolor="#1e293b",
            plot_bgcolor="#0f172a",
            hovermode="x unified",
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
            margin=dict(l=40, r=20, t=50, b=40),
            yaxis=dict(title="괴리율 (%)", gridcolor="#334155"),
            xaxis=dict(gridcolor="#334155")
        )
        st.plotly_chart(fig2, use_container_width=True)

    # [차트 3 및 배당 분석 요약]: 배당금 및 배당수익률 비교
    col_div1, col_div2 = st.columns([6, 6])
    
    with col_div1:
        # 배당금 & 배당수익률 비교 바 차트
        fig3 = make_subplots(specs=[[{"secondary_y": True}]])
        
        categories = ["보통주", "우선주"]
        dps_values = [target_row["보통주 배당금"], target_row["우선주 배당금"]]
        yield_values = [target_row["보통주 배당수익률(A)"], target_row["우선주 배당수익률(B)"]]

        # 배당금(원) Bar
        fig3.add_trace(
            go.Bar(
                x=categories,
                y=dps_values,
                name="주당 배당금(원)",
                marker_color=["#60a5fa", "#f472b6"],
                text=[f"{v:,.0f}원" for v in dps_values],
                textposition="auto"
            ),
            secondary_y=False
        )

        # 배당수익률(%) Scatter/Line
        fig3.add_trace(
            go.Scatter(
                x=categories,
                y=yield_values,
                name="배당수익률(%)",
                mode="markers+text",
                marker=dict(size=14, color="#fbbf24"),
                text=[f"{v:.2f}%" for v in yield_values],
                textposition="top center"
            ),
            secondary_y=True
        )

        fig3.update_layout(
            title=dict(
                text=f"<b>보통주 vs 우선주 배당 비교</b>",
                font=dict(color="#f8fafc", size=15)
            ),
            template="plotly_dark",
            paper_bgcolor="#1e293b",
            plot_bgcolor="#0f172a",
            margin=dict(l=40, r=40, t=50, b=40),
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
        )
        fig3.update_yaxes(title_text="주당 배당금 (원)", secondary_y=False, gridcolor="#334155")
        fig3.update_yaxes(title_text="배당수익률 (%)", secondary_y=True, gridcolor="#334155")

        st.plotly_chart(fig3, use_container_width=True)

    with col_div2:
        # 우선주 투자 매력도 및 체크포인트 카드
        current_disc = target_row['괴리율(%)']
        div_premium = target_row['우선주 배당수익률(B)'] - target_row['보통주 배당수익률(A)']
        hist_avg = df_hist['괴리율'].mean()
        diff_vs_avg = current_disc - hist_avg

        with st.container(border=True):
            st.markdown("<div style='font-size: 1.15rem; font-weight: 700; color: #8AB4F8; margin-bottom: 12px;'>💡 투자 매력도 & 핵심 체크포인트</div>", unsafe_allow_html=True)

            # 1. 괴리율 진단
            st.markdown("**1. 괴리율(할인율) 진단:**")
            if diff_vs_avg > 2.0:
                st.markdown(f":green[**과거 평균 대비 고괴리(저평가) 구간 (+{diff_vs_avg:.1f}%p)**]")
                st.caption("현재 보통주 대비 우선주의 할인 폭이 과거 평균보다 커서 가격 메리트가 높은 매력적인 구간입니다.")
            elif diff_vs_avg < -2.0:
                st.markdown(f":red[**과거 평균 대비 저괴리(축소) 구간 ({diff_vs_avg:.1f}%p)**]")
                st.caption("현재 괴리율이 과거 평균보다 좁혀져 있어 보통주와의 주가 갭이 상대적으로 작은 구간입니다.")
            else:
                st.markdown(f":blue[**과거 평균 수준의 적정 괴리 구간 ({diff_vs_avg:+.1f}%p)**]")
                st.caption("현재 괴리율이 과거 평균 범위 내에서 안정적으로 유지되고 있습니다.")

            st.divider()

            # 2. 배당 프리미엄
            st.markdown("**2. 배당 프리미엄:**")
            st.markdown(f"<span style='font-size: 1.3rem; font-weight: 800; color: #fbbf24;'>+{div_premium:.2f}%p</span>", unsafe_allow_html=True)
            st.caption(f"동일 금액 투자 시 우선주가 보통주보다 연간 약 **{div_premium:.2f}%p** 더 높은 배당 수익을 제공합니다.")

            st.divider()

            # 3. 투자 전략 제언
            st.markdown("**3. 투자 전략 제언:**")
            st.caption("우선주는 의결권이 없는 대신 **높은 배당수익률**과 **괴리율 축소 시 자본차익**을 동시에 노릴 수 있습니다. 시가총액이 크고 배당성향이 안정적인 대형 우선주 중심으로 접근하는 것이 유동성 리스크를 줄이는 데 유리합니다.")
